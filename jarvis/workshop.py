from __future__ import annotations

from pathlib import Path
import re
from xml.sax.saxutils import escape

_PROJECT_ROOT = Path(__file__).parent.parent
_WORKSPACE = _PROJECT_ROOT / "workspace"


def workspace_path() -> Path:
    _WORKSPACE.mkdir(parents=True, exist_ok=True)
    return _WORKSPACE


def _safe_stem(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9ÄÖÜäöüß _.-]+", "", value).strip(" .")
    value = re.sub(r"\s+", " ", value)
    return (value or "Dokument")[:100]


def _unique_path(stem: str, suffix: str) -> Path:
    root = workspace_path()
    candidate = root / f"{_safe_stem(stem)}{suffix}"
    index = 2
    while candidate.exists():
        candidate = root / f"{_safe_stem(stem)}-{index}{suffix}"
        index += 1
    return candidate


def _resolve_existing(name: str) -> Path:
    root = workspace_path().resolve()
    candidate = (root / name).resolve()
    if root != candidate and root not in candidate.parents:
        raise ValueError("Pfad liegt außerhalb des COMPUTER Workspace.")
    if not candidate.exists() or not candidate.is_file():
        raise FileNotFoundError(name)
    return candidate


def create_document(title: str, content: str, file_format: str = "docx") -> str:
    """Create a new draft document without overwriting existing files."""
    fmt = file_format.lower().lstrip(".")
    title = title.strip() or "Dokument"

    if fmt == "txt":
        path = _unique_path(title, ".txt")
        path.write_text(content, encoding="utf-8")
    elif fmt == "docx":
        from docx import Document

        path = _unique_path(title, ".docx")
        doc = Document()
        doc.add_heading(title, level=1)
        for paragraph in re.split(r"\n\s*\n", content.strip()):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
        doc.save(path)
    elif fmt == "pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        path = _unique_path(title, ".pdf")
        styles = getSampleStyleSheet()
        story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 14)]
        for paragraph in re.split(r"\n\s*\n", content.strip()):
            if paragraph.strip():
                story.append(Paragraph(escape(paragraph.strip()).replace("\n", "<br/>"), styles["BodyText"]))
                story.append(Spacer(1, 9))
        SimpleDocTemplate(str(path), pagesize=A4, title=title).build(story)
    else:
        raise ValueError("Unterstützte Formate: docx, pdf, txt.")

    return f"Dokument erstellt: {path.name} | {path}"


def read_document(name: str) -> str:
    """Read TXT, DOCX or text from a PDF in the COMPUTER workspace."""
    path = _resolve_existing(name)
    suffix = path.suffix.lower()

    if suffix == ".txt":
        text = path.read_text(encoding="utf-8")
    elif suffix == ".docx":
        from docx import Document

        doc = Document(path)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(lines)
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    else:
        raise ValueError("Unterstützte Formate: docx, pdf, txt.")

    return text[:30000]


def revise_document(name: str, find_text: str, replacement: str) -> str:
    """Create a revised copy. Existing originals are never overwritten."""
    if not find_text:
        raise ValueError("Suchtext darf nicht leer sein.")

    path = _resolve_existing(name)
    suffix = path.suffix.lower()
    revised_title = f"{path.stem}-revised"

    if suffix == ".txt":
        original = path.read_text(encoding="utf-8")
        if find_text not in original:
            return f"Suchtext nicht gefunden: {find_text}"
        revised = original.replace(find_text, replacement)
        out = _unique_path(revised_title, ".txt")
        out.write_text(revised, encoding="utf-8")
        return f"Überarbeitete Kopie erstellt: {out.name}"

    if suffix == ".docx":
        from docx import Document

        doc = Document(path)
        hits = 0
        for paragraph in doc.paragraphs:
            if find_text in paragraph.text:
                for run in paragraph.runs:
                    if find_text in run.text:
                        run.text = run.text.replace(find_text, replacement)
                        hits += 1
                if find_text in paragraph.text:
                    paragraph.text = paragraph.text.replace(find_text, replacement)
                    hits += 1
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if find_text in cell.text:
                        cell.text = cell.text.replace(find_text, replacement)
                        hits += 1
        if hits == 0:
            return f"Suchtext nicht gefunden: {find_text}"
        out = _unique_path(revised_title, ".docx")
        doc.save(out)
        return f"Überarbeitete Kopie erstellt: {out.name}"

    if suffix == ".pdf":
        # PDF text is not safely editable in-place. Rebuild a revised copy from
        # extracted text instead of pretending layout-preserving edits occurred.
        original = read_document(path.name)
        if find_text not in original:
            return f"Suchtext nicht gefunden: {find_text}"
        revised = original.replace(find_text, replacement)
        result = create_document(revised_title, revised, "pdf")
        return result + " | Hinweis: PDF wurde aus extrahiertem Text neu aufgebaut; Originallayout bleibt nicht garantiert."

    raise ValueError("Unterstützte Formate: docx, pdf, txt.")


def list_documents() -> str:
    files = sorted(p for p in workspace_path().iterdir() if p.is_file())
    if not files:
        return "COMPUTER Workspace ist leer."
    return "\n".join(f"{p.name} ({p.stat().st_size // 1024} KB)" for p in files[:100])
